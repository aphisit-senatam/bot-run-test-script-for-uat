from selenium import webdriver 
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.common.by import By
import time
import pytesseract
from PIL import Image
import re
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
import os
from openpyxl import load_workbook
from selenium.webdriver.common.action_chains import ActionChains
from PIL import Image, ImageEnhance
options = webdriver.ChromeOptions()
options.add_argument("--start-maximized")
wb = load_workbook(r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inboundDB.xlsx")
ws = wb["Inbound"]
DockNo = ws["B9"].value
ZoneGroup = ws["B10"].value
EquipmentType = ws["B11"].value
VehicleType = ws["B12"].value
ReceivingLocation = ws["B13"].value
DockCapacity = ws["B14"].value
OpenTimeFM1 = ws["B15"].value
OpenTimeTO1 = ws["B16"].value
OpenTimeFM2 = ws["B17"].value
OpenTimeTO2 = ws["B18"].value
OpenTimeFM3 = ws["B19"].value
OpenTimeTO3 = ws["B20"].value

wb = load_workbook(r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inboundDB.xlsx")
login = wb["URL WMS"]
url_uat = login["B2"].value
user = login["B3"].value
password = login["B4"].value



driver = webdriver.Chrome(
    service=Service(ChromeDriverManager().install()),
    options=options
)
driver.get(url_uat)
driver.find_element(By.ID,"details-button").click()
driver.find_element(By.ID,"proceed-link").click()
driver.find_element(By.ID,"CURUSERID").send_keys(user)
driver.find_element(By.ID,"CURURPD").send_keys(password)
time.sleep(1)
len_cap = 0
loop_count = 0
MAX_LOOP = 10

while loop_count < MAX_LOOP:
    print(f"\n=== LOOP {loop_count + 1} ===")

    try:
        img = driver.find_element(By.ID, "checknumImg")
        time.sleep(0.5)
        img.screenshot(
            r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\capthca\capthca.png"
        )

        image = Image.open(
            r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\capthca\capthca.png"
        )
        image = image.convert("L")
        image = ImageEnhance.Contrast(image).enhance(3.0)
        image = image.point(lambda x: 0 if x < 140 else 255, '1')

        captcha = pytesseract.image_to_string(
            image,
            config="--psm 7 --oem 1 -c tessedit_char_whitelist=0123456789"
        )

        Revisecaptcha = re.sub(r"\D", "", captcha)
        len_cap = len(Revisecaptcha)

        print("captcha OCR =", Revisecaptcha, "| len =", len_cap)

        if len_cap != 4:
            print("OCR not 4 ")
            driver.execute_script("arguments[0].click();", img)
            loop_count += 1
            continue

        inputbox = driver.find_element(By.ID, "CHECKNUM")
        inputbox.clear()
        inputbox.send_keys(Revisecaptcha)

        driver.find_element(By.ID, "submit").click()
        time.sleep(1)


        errors = driver.find_elements(By.ID, "login_validationMessage")

        if errors and errors[0].is_displayed():
            print("wrong captcha ")
            driver.execute_script("arguments[0].click();", img)
            loop_count += 1
            continue

        print("captcha passing ")
        break

    except Exception as e:
        print("ERROR Loop:", e)
        loop_count += 1

if loop_count >= MAX_LOOP:
    print("loop is max try again")
    driver.quit()

btn = driver.find_element(By.XPATH, "//input[@value='Force Login']")
driver.execute_script("arguments[0].click();", btn)
#funtion 

def dock():
    function = 'A1010 Dock - 1010'
    buttonnv = driver.find_element(By.ID, "menu-btn").click()
    #time.sleep(3)
    box = driver.find_element(By.XPATH,"//input[@placeholder='Function Retrieval']")
    box.clear()
    time.sleep(2)
    box.send_keys(function + Keys.ENTER)
    time.sleep(2)
    add = driver.find_element(By.XPATH, "//div[@class='dhxtoolbar_text' and text()='Add']").click()
    time.sleep(1)
    action = ActionChains(driver)
    action.send_keys(str(DockNo))
    action.send_keys(Keys.TAB)
    action.send_keys(str(ZoneGroup))
    action.send_keys(Keys.TAB)
    action.send_keys(str(EquipmentType))
    action.send_keys(Keys.TAB)
    action.send_keys(str(VehicleType))
    action.send_keys(Keys.TAB)
    action.send_keys(str(ReceivingLocation))
    action.send_keys(Keys.TAB)
    action.send_keys(str(DockCapacity))
    action.perform()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ1\1สามารถAddได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    save = driver.find_element(By.XPATH, "//button[.//i[contains(@class,'icon-save')]]").click()
    time.sleep(1)
    dockOpenTime = driver.find_element(By.XPATH,"//div[contains(@class,'dhxtabbar_tab_text') and normalize-space()='Dock Open Time']")
    driver.execute_script("arguments[0].click();", dockOpenTime)
    time.sleep(1)
    wait = WebDriverWait(driver, 20)
    add_btn = wait.until(EC.presence_of_element_located(( By.XPATH,"//div[@id='A1010S01_DOCKOPENTIME_cell']""//div[contains(@class,'dhx_toolbar_btn') and .//div[normalize-space()='Add']]")))
    add_btn.click()
    time.sleep(1)
    Opentime= driver.find_element(By.NAME, "openTimeFM")
    Opentime.clear
    Opentime.send_keys(OpenTimeFM1)
    Opentime.send_keys(Keys.TAB)
    OpentimeTo = ActionChains(driver)
    OpentimeTo.send_keys(OpenTimeTO1)
    OpentimeTo.send_keys(Keys.TAB)
    OpentimeTo.send_keys(Keys.TAB)
    OpentimeTo.send_keys(Keys.ENTER)
    OpentimeTo.perform()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ1\2Addเวลาประตูได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    #ข้อ2
    time.sleep(1)
    wait = WebDriverWait(driver, 20)
    add_btn = wait.until(EC.presence_of_element_located(( By.XPATH,"//div[@id='A1010S01_DOCKOPENTIME_cell']""//div[contains(@class,'dhx_toolbar_btn') and .//div[normalize-space()='Add']]")))
    add_btn.click()
    time.sleep(1)
    Opentime= driver.find_element(By.NAME, "openTimeFM")
    Opentime.clear
    Opentime.send_keys(OpenTimeFM2)
    Opentime.send_keys(Keys.TAB)
    OpentimeTo = ActionChains(driver)
    OpentimeTo.send_keys(OpenTimeTO2)
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ2\1สามารถAddเวลาประตูเพิ่มได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    OpentimeTo.send_keys(Keys.TAB)
    OpentimeTo.send_keys(Keys.TAB)
    OpentimeTo.send_keys(Keys.ENTER)
    OpentimeTo.perform()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ2\2Addเวลาประตูเพิ่มได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    #ข้อ3
    time.sleep(1)
    Opentime = wait.until(EC.element_to_be_clickable((By.NAME, "openTimeFM")))
    Opentime.clear()
    Opentime.send_keys(OpenTimeFM3)
    OpentimeTo = wait.until(EC.element_to_be_clickable((By.NAME, "openTimeTO")))
    OpentimeTo.clear()
    OpentimeTo.send_keys(OpenTimeTO3)
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ3\1สามารถแก้ไขได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    savefixtime = ActionChains(driver)
    savefixtime.send_keys(Keys.TAB)
    savefixtime.send_keys(Keys.TAB)
    savefixtime.send_keys(Keys.ENTER)
    savefixtime.perform()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ3\2แก้ไขได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    #ข้อ4
    time.sleep(1)
    wait = WebDriverWait(driver, 20)
    delete_btn = wait.until(EC.presence_of_element_located(( By.XPATH,"//div[@id='A1010S01_DOCKOPENTIME_cell']""//div[contains(@class,'dhx_toolbar_btn') and .//div[normalize-space()='Delete']]")))
    delete_btn.click()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ4\1สามารถลบได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    confirm_btn = wait.until(EC.element_to_be_clickable((By.ID,"layer-index-btn--0"))).click()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ4\2ลบได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    elems = WebDriverWait(driver, 10).until(EC.presence_of_all_elements_located((By.XPATH, '//*[@id="a_cell"]/div[4]'))) #ปิดแทบ Detail
    elems[0].click()

    #ข้อ5
    #เปิดแทบ serch
    header = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH,"//div[contains(@class,'dhx_cell_hdr_text') and .//span[text()='Query Condition']]")))
    ActionChains(driver).move_to_element(header).click().perform() 
    #serching input dock
    wait = WebDriverWait(driver, 20)
    dock = wait.until( EC.element_to_be_clickable((By.NAME, "DOCK")))
    dock.clear()
    dock.send_keys("DOCKAA006")
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\1ค้นหาDockNo.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    wait = WebDriverWait(driver, 20)
    #ปุ่ม serch
    search_btn = wait.until(EC.element_to_be_clickable((By.NAME, "A1010searchForHeader_searchBtn")))
    search_btn.click()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\2ค้นหาDockNoได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
  #serching combo list zone group
    header = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH,"//div[contains(@class,'dhx_cell_hdr_text') and .//span[text()='Query Condition']]")))
    ActionChains(driver).move_to_element(header).click().perform()
    dock.clear()
    time.sleep(1)
    dock = wait.until(EC.element_to_be_clickable((By.NAME, "DOCK")))
    dock.click()  # สำคัญ: ให้ focus อยู่ที่ field นี้ก่อน
    zonegouplist = ActionChains(driver)
    zonegouplist.send_keys(Keys.TAB)
    zonegouplist.send_keys("Ambient Warehouse")
    zonegouplist.send_keys(Keys.ENTER)
    zonegouplist.send_keys(Keys.ENTER)
    zonegouplist.perform()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\3ค้นหาzonegroupได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
  #clear combo list
    time.sleep(1)
    header = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH,"//div[contains(@class,'dhx_cell_hdr_text') and .//span[text()='Query Condition']]")))
    ActionChains(driver).move_to_element(header).click().perform()
    dock = wait.until(EC.element_to_be_clickable((By.NAME, "DOCK")))
    dock.click()  
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\3ค้นหาzonegroup.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    actionlistserch = ActionChains(driver)
    actionlistserch.send_keys(Keys.TAB)                
    actionlistserch.key_down(Keys.CONTROL).send_keys("a").key_up(Keys.CONTROL)
    actionlistserch.send_keys(Keys.DELETE)             #delete combo list
    actionlistserch.send_keys(Keys.TAB) 
    actionlistserch.send_keys(Keys.ARROW_DOWN) 
    actionlistserch.send_keys(Keys.ARROW_DOWN) 
    actionlistserch.send_keys(Keys.ENTER) 
    actionlistserch.send_keys(Keys.ENTER)  #INFLAG SERCH
    actionlistserch.perform()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\5ค้นหาInFlagได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    time.sleep(1)
    header = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH,"//div[contains(@class,'dhx_cell_hdr_text') and .//span[text()='Query Condition']]")))
    ActionChains(driver).move_to_element(header).click().perform()
    dock = wait.until(EC.element_to_be_clickable((By.NAME, "DOCK")))
    dock.click()  
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\4ค้นหาInFlag.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    actionlistserch2 = ActionChains(driver)
    actionlistserch2.send_keys(Keys.TAB) 
    actionlistserch2.send_keys(Keys.TAB) 
    actionlistserch2.send_keys(Keys.DELETE)             #delete Inflag
    actionlistserch2.send_keys(Keys.TAB) 
    actionlistserch2.send_keys(Keys.ARROW_DOWN) 
    actionlistserch2.send_keys(Keys.ARROW_DOWN) 
    actionlistserch2.send_keys(Keys.ENTER) 
    actionlistserch2.send_keys(Keys.ENTER)  #SO FLAG SERCH
    actionlistserch2.perform()
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\6ค้นหาSOFlagได้.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
    time.sleep(1)
    header = WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.XPATH,"//div[contains(@class,'dhx_cell_hdr_text') and .//span[text()='Query Condition']]")))
    ActionChains(driver).move_to_element(header).click().perform()
    dock = wait.until(EC.element_to_be_clickable((By.NAME, "DOCK")))
    dock.click()  
    time.sleep(1)
    filename = r"C:\Users\aphisit.sen\Desktop\งานโปรเจค\DC5 Project\selenuim bot project UNT\inbound\Dock\ข้อ5\6ค้นหาSOFlag.png"
    folder = os.path.dirname(filename)
    os.makedirs(folder, exist_ok=True)
    driver.save_screenshot(filename)
    print("screenshot save...")
   
#end of function
time.sleep(2)
print('after',driver.current_url)
WebDriverWait(driver, 60).until(lambda d: "#/home" in d.current_url)
print('before',driver.current_url)
print("running... function 1")
dock()
print("function 1 Done")
time.sleep(1)
print("****ทำงานจบเรียบร้อยครับ Code Selenuim By Ahisit Senatam(Aof)****")
driver.quit()